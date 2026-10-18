"""Test harness: real PostgreSQL (TEST_DATABASE_URL), schema built by Alembic migrations,
a controllable clock, and helpers. All customer-like data is synthetic."""

from __future__ import annotations

import io
import os
import warnings
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from sqlalchemy import create_engine, text

warnings.filterwarnings("ignore", message=".*httpx2.*")

from alembic import command  # noqa: E402
from alembic.config import Config  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from routewise_api.config import load_settings  # noqa: E402
from routewise_api.main import create_app  # noqa: E402
from routewise_api.models import Rider, User  # noqa: E402
from routewise_api.security import hash_password  # noqa: E402

API_DIR = Path(__file__).resolve().parents[1]
ORIGIN = "http://localhost:5173"
TABLES = ("plan_orders, plan_routes, route_plans, planning_runs, rider_endpoints, round_riders, depots, assignments, "
          "orders, import_batches, delivery_rounds, sessions, login_failures, riders, users")


def _test_url() -> str:
    load_settings()  # loads the git-ignored .env into os.environ
    url = os.environ.get("TEST_DATABASE_URL")
    if not url:
        pytest.exit("TEST_DATABASE_URL is not set; run scripts/dev-db.sh init", returncode=2)
    return url


def alembic_config(url: str) -> Config:
    cfg = Config(str(API_DIR / "alembic.ini"))
    cfg.set_main_option("script_location", str(API_DIR / "migrations"))
    cfg.cmd_opts = type("o", (), {"x": [f"url={url}"]})()
    return cfg


def reset_schema(url: str) -> None:
    eng = create_engine(url)
    with eng.begin() as conn:
        conn.execute(text("DROP SCHEMA public CASCADE; CREATE SCHEMA public"))
    eng.dispose()


@pytest.fixture(scope="session")
def db_url() -> str:
    url = _test_url()
    reset_schema(url)
    command.upgrade(alembic_config(url), "head")
    return url


class Clock:
    def __init__(self):
        self.now = datetime(2026, 10, 13, 9, 0, tzinfo=timezone.utc)

    def __call__(self) -> datetime:
        return self.now

    def advance(self, **kw) -> None:
        self.now += timedelta(**kw)


@pytest.fixture
def clock() -> Clock:
    return Clock()


@pytest.fixture
def settings(db_url):
    s = load_settings(database_url=db_url)
    return replace(s, resolve_short_links=False)  # tests never touch the network


@pytest.fixture
def app(settings, clock):
    application = create_app(settings, clock=clock)
    with application.state.engine.begin() as conn:
        conn.execute(text(f"TRUNCATE {TABLES} RESTART IDENTITY CASCADE"))
        conn.execute(text("ALTER SEQUENCE order_code_seq RESTART WITH 1001"))
    yield application
    application.state.planning.shutdown()
    application.state.engine.dispose()


@pytest.fixture
def db(app):
    s = app.state.session_factory()
    yield s
    s.close()


def make_user(app, username: str, role: str, password: str = "correct horse battery", *,
              rider_name: str | None = None, vehicle: str | None = None) -> User:
    with app.state.session_factory() as s:
        user = User(username=username, role=role, password_hash=hash_password(password))
        s.add(user)
        s.flush()
        if role == "rider":
            s.add(Rider(user_id=user.id, display_name=rider_name or username, vehicle=vehicle))
        s.commit()
        s.refresh(user)
        return user


class Browser:
    """A TestClient with its own cookie jar plus the CSRF token of its session."""

    def __init__(self, app, base_url: str = "http://testserver"):
        self.client = TestClient(app, base_url=base_url)
        self.csrf: str | None = None

    def login(self, username: str, password: str = "correct horse battery", origin: str = ORIGIN):
        r = self.client.post("/api/auth/login", json={"username": username, "password": password},
                             headers={"Origin": origin})
        if r.status_code == 200:
            self.csrf = self.client.get("/api/auth/session").json()["csrf_token"]
        return r

    def headers(self, csrf: bool = True, origin: str | None = ORIGIN) -> dict:
        h = {}
        if origin:
            h["Origin"] = origin
        if csrf and self.csrf:
            h["X-CSRF-Token"] = self.csrf
        return h

    def get(self, url, **kw):
        return self.client.get(url, **kw)

    def send(self, method: str, url: str, *, csrf: bool = True, origin: str | None = ORIGIN, **kw):
        headers = {**self.headers(csrf, origin), **kw.pop("headers", {})}
        return self.client.request(method, url, headers=headers, **kw)


@pytest.fixture
def browser(app):
    def _make(**kw) -> Browser:
        return Browser(app, **kw)

    return _make


@pytest.fixture
def admin(app, browser) -> Browser:
    make_user(app, "admin1", "admin")
    b = browser()
    assert b.login("admin1").status_code == 200
    return b


def xlsx(rows: list[list], *, sheet_title: str = "Orders", number_formats: dict | None = None,
         extra_sheet: bool = False) -> bytes:
    """Build a synthetic workbook. number_formats maps (row, col) 1-based -> Excel format."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r in rows:
        ws.append(r)
    for (r, c), fmt in (number_formats or {}).items():
        ws.cell(row=r, column=c).number_format = fmt
    if extra_sheet:
        wb.create_sheet("Notes").append(["ignored"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["Name", "Product", "Contact", "Location", "Delivery Instructions"]
PIN = "https://www.google.com/maps/place/X/@28.6800,77.2000,17z/data=!3d28.6862643!4d77.2077300"


def upload(b: Browser, data: bytes, filename: str = "orders.xlsx"):
    return b.send("POST", "/api/admin/imports",
                  files={"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
